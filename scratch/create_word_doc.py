import os
from pathlib import Path
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def create_element(name):
    return OxmlElement(name)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_heading_styled(doc, text, level):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.bold = True
    if level == 1:
        run.font.size = Pt(18)
        run.font.color.rgb = RGBColor(16, 124, 65) # Dark Emerald Green
    elif level == 2:
        run.font.size = Pt(14)
        run.font.color.rgb = RGBColor(33, 37, 41)
    elif level == 3:
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(100, 100, 100)
    return p

def main():
    handover_dir = Path(r"C:\Users\RICHIE\Documents\AGROWATCH-1\AgroWatch_APK_Handover")
    doc_path = handover_dir / "AgroWatch_APK_Handover.docx"

    doc = Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Title Block
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_after = Pt(4)
    r_title = title_p.add_run("AGROWATCH MOBILE APPLICATION HANDOVER DOCUMENTATION")
    r_title.font.name = 'Calibri'
    r_title.font.size = Pt(22)
    r_title.bold = True
    r_title.font.color.rgb = RGBColor(16, 124, 65)

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_p.paragraph_format.space_after = Pt(18)
    r_sub = sub_p.add_run("APK Release v1.0, Technical Specifications, API Integration & Verification Report")
    r_sub.font.name = 'Calibri'
    r_sub.font.size = Pt(12)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(80, 80, 80)

    # Metadata Box Table
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Application Name & Package:", "AgroWatch (agrowatch.app.com)"),
        ("Release Artifact:", "Agrowatch_v1.0.apk (Version 1.0, Code 1)"),
        ("Backend Cloud API:", "https://agrowatch-5h12.onrender.com/api"),
        ("Date & Quality Assurance:", "May 11, 2026 | QA Verification: PASSED")
    ]
    for idx, (label, val) in enumerate(meta_data):
        row = meta_table.rows[idx]
        cell_lbl, cell_val = row.cells[0], row.cells[1]

        cell_lbl.width = Inches(2.3)
        cell_val.width = Inches(4.5)

        set_cell_background(cell_lbl, "F0F4F1")
        set_cell_background(cell_val, "FAFAFA")
        set_cell_margins(cell_lbl, top=80, bottom=80, left=100, right=100)
        set_cell_margins(cell_val, top=80, bottom=80, left=100, right=100)

        p_l = cell_lbl.paragraphs[0]
        r_l = p_l.add_run(label)
        r_l.bold = True
        r_l.font.size = Pt(10)

        p_v = cell_val.paragraphs[0]
        r_v = p_v.add_run(val)
        r_v.font.size = Pt(10)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # --- SECTION 1: INSTALLATION & USER GUIDE ---
    add_heading_styled(doc, "1. Installation & User Setup Guide", level=1)

    p = doc.add_paragraph()
    p.add_run("Prerequisites: ").bold = True
    p.add_run("Android 7.0 (API level 24) or higher, 100 MB free storage, and an active internet connection for live cloud API synchronization and AI model inference.")

    add_heading_styled(doc, "Step-by-Step Installation", level=2)
    steps = [
        "Transfer 'Agrowatch_v1.0.apk' to the Android target device via USB, WhatsApp, or cloud storage.",
        "Open File Manager on the device and locate 'Agrowatch_v1.0.apk'.",
        "Tap the APK to begin installation. If prompted to allow unknown sources, enable 'Install Unknown Apps' for File Manager in Settings -> Security.",
        "Tap 'Install' and launch AgroWatch from the app drawer upon completion.",
        "Grant Camera & Storage permissions when prompted for crop scanning features."
    ]
    for s in steps:
        doc.add_paragraph(s, style='List Bullet')

    add_heading_styled(doc, "Demo Accounts for Testing & Review", level=2)
    demo_table = doc.add_table(rows=4, cols=4)
    demo_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Role", "Username / Phone", "Password", "Permissions & Access Scope"]
    hdr_row = demo_table.rows[0]
    for i, h in enumerate(headers):
        cell = hdr_row.cells[i]
        set_cell_background(cell, "107C41")
        set_cell_margins(cell, top=100, bottom=100, left=100, right=100)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(10)

    demo_rows = [
        ("Farmer", "farmer_demo", "Password123!", "Crop management, Drone Scan upload, Market Produce Listing"),
        ("Buyer", "buyer_demo", "Password123!", "Marketplace browsing, direct farmer messaging, produce purchase"),
        ("System Admin", "admin_demo", "AdminPassword123!", "System-wide metrics, scan archive audit, user administration")
    ]
    for idx, row_data in enumerate(demo_rows, start=1):
        row = demo_table.rows[idx]
        bg = "F9F9F9" if idx % 2 == 1 else "FFFFFF"
        for j, val in enumerate(row_data):
            cell = row.cells[j]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(9.5)
            if j == 0:
                r.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # --- SECTION 2: SCREENSHOTS & UI FIGURES ---
    add_heading_styled(doc, "2. User Interface Screenshots & Figures", level=1)

    screenshots_info = [
        ("fig4_1_home.png", "Figure 4.1: AgroWatch Home & Overview Dashboard", "Displays key platform stats, quick navigation actions, and recent activity overview."),
        ("fig4_2_login.png", "Figure 4.2: Authentication & Login Portal", "User login interface supporting Farmer, Buyer, and Administrator role-based access."),
        ("fig4_3_scan.png", "Figure 4.3: Drone Crop Health Scan Interface", "Camera & image upload portal for initiating crop health scans."),
        ("fig4_4_diagnosis.png", "Figure 4.4: AI Disease Diagnosis Report", "Presents YOLOv8 bounding box detections, confidence scores, and expert treatment advice."),
        ("fig4_5_market.png", "Figure 4.5: Produce Market Exchange", "Direct agricultural marketplace connecting verified farmers with produce buyers."),
        ("fig4_6_farms.png", "Figure 4.6: Farm Plots Management Registry", "Field plot registration and geographic farm tracking across agricultural regions.")
    ]

    screenshots_dir = handover_dir / "screenshots"
    for img_name, fig_title, fig_desc in screenshots_info:
        img_path = screenshots_dir / img_name
        if img_path.exists():
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(8)
            p_img.paragraph_format.space_after = Pt(2)

            # Width constrained for clean page layout
            p_img.add_run().add_picture(str(img_path), width=Inches(3.8))

            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_after = Pt(14)
            p_cap.paragraph_format.keep_with_next = False

            r_cap = p_cap.add_run(fig_title + " – ")
            r_cap.bold = True
            r_cap.font.size = Pt(9.5)
            r_cap.font.color.rgb = RGBColor(16, 124, 65)

            r_desc = p_cap.add_run(fig_desc)
            r_desc.font.size = Pt(9.5)
            r_desc.font.italic = True

    # --- SECTION 3: BACKEND API & ML MODEL INTEGRATION ---
    add_heading_styled(doc, "3. Backend API & ML Model Integration Architecture", level=1)

    p_api = doc.add_paragraph()
    p_api.add_run("The AgroWatch mobile application communicates with a Django REST Framework backend hosted on Render Cloud (").font.size = Pt(10)
    p_api.add_run("https://agrowatch-5h12.onrender.com/api").bold = True
    p_api.add_run("). All requests utilize token authentication and structured JSON payloads.").font.size = Pt(10)

    add_heading_styled(doc, "Crop-Specific ML Model Mapping", level=2)

    ml_table = doc.add_table(rows=4, cols=4)
    ml_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    ml_headers = ["Crop Species", "Model Weights File", "Inference Engine", "Detected Diseases & Conditions"]
    for i, h in enumerate(ml_headers):
        cell = ml_table.rows[0].cells[i]
        set_cell_background(cell, "107C41")
        set_cell_margins(cell, top=100, bottom=100, left=100, right=100)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(10)

    ml_rows = [
        ("Tomato", "mlweights/tomato.pt", "YOLOv8 Custom", "Tomato Blight, Yellow Leaf Curl, Spot, Healthy"),
        ("Maize", "mlweights/maize.pt", "YOLOv8 Custom", "Blight, Common Rust, Gray Leaf Spot, Healthy"),
        ("Pineapple", "mlweights/pineapple.pt", "YOLOv8 Custom", "Heart Rot, Mealybug Wilt, Black Rot, Healthy")
    ]
    for idx, row_data in enumerate(ml_rows, start=1):
        row = ml_table.rows[idx]
        bg = "F9F9F9" if idx % 2 == 1 else "FFFFFF"
        for j, val in enumerate(row_data):
            cell = row.cells[j]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(9.5)
            if j == 0 or j == 1:
                r.bold = True

    add_heading_styled(doc, "Pre-Inference Crop & Image Validation (crop_validator.py)", level=2)
    p_val = doc.add_paragraph()
    p_val.add_run("To ensure high diagnostic accuracy and prevent model degradation, uploaded imagery undergoes multi-stage pre-inference verification before YOLOv8 execution:")

    val_points = [
        "Non-Agricultural Filter: Automatically detects and rejects human faces, skin tones, vehicles, indoor background items, or corrupted files.",
        "Species Cross-Validation: Verifies leaf/field morphology against selected crop type (e.g. rejects Tomato image on Maize scan).",
        "Cloud-Safe Architecture: Operates via lightweight OpenCV + NumPy with MobileNet_V3_Small fallback."
    ]
    for vp in val_points:
        doc.add_paragraph(vp, style='List Bullet')

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # --- SECTION 4: COMPREHENSIVE QA TEST LOG ---
    add_heading_styled(doc, "4. QA Verification & Test Execution Log", level=1)

    test_table = doc.add_table(rows=13, cols=5)
    test_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    test_headers = ["Test ID", "Category", "Description", "Status", "Verification Notes"]
    for i, h in enumerate(test_headers):
        cell = test_table.rows[0].cells[i]
        set_cell_background(cell, "107C41")
        set_cell_margins(cell, top=100, bottom=100, left=100, right=100)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(9.5)

    test_data = [
        ("TC-001", "Installation", "Install APK on Android 13/14 device", "PASS", "Installed cleanly without package conflict."),
        ("TC-002", "App Startup", "Launch app & check splash screen & layout", "PASS", "Splash screen displays; adaptive edge-to-edge fit."),
        ("TC-003", "Authentication", "Login with demo accounts (Farmer/Buyer/Admin)", "PASS", "Tokens stored securely; correct role routing."),
        ("TC-004", "Registration", "Create new user account via registration form", "PASS", "New user registered and authenticated."),
        ("TC-005", "Farm Plot Reg", "Register new farm (Tomato, 2.5ha, Ashanti)", "PASS", "Farm object created & listed in Farms tab."),
        ("TC-006", "Scan Upload", "Upload crop leaf image for disease scan", "PASS", "Image sent to /api/scans/; YOLO model invoked."),
        ("TC-007", "Crop Validation", "Upload invalid image (non-crop) for scan", "PASS", "crop_validator rejected non-agricultural image."),
        ("TC-008", "AI Diagnosis", "View disease detection report & expert tips", "PASS", "Blighting detected with treatment advice rendered."),
        ("TC-009", "Marketplace", "Create produce listing & view market cards", "PASS", "Listing visible to buyers; contact details linked."),
        ("TC-010", "Messaging", "Send inquiry message between Buyer and Farmer", "PASS", "Message sent and received via API endpoints."),
        ("TC-011", "Admin Portal", "Log in as admin_demo and check system metrics", "PASS", "System stats, user list & scans displayed."),
        ("TC-012", "Network Error", "Test offline / poor network behavior", "PASS", "Graceful error toast notifications displayed.")
    ]

    for idx, row_data in enumerate(test_data, start=1):
        row = test_table.rows[idx]
        bg = "F9F9F9" if idx % 2 == 1 else "FFFFFF"
        for j, val in enumerate(row_data):
            cell = row.cells[j]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(9)
            if j == 0:
                r.bold = True
            elif j == 3:
                r.bold = True
                r.font.color.rgb = RGBColor(16, 124, 65) # Green PASS

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # --- SECTION 5: VERSION & TECH STACK ---
    add_heading_styled(doc, "5. Version Details & Technology Stack", level=1)

    tech_table = doc.add_table(rows=6, cols=2)
    tech_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tech_data = [
        ("Application Name & Version", "AgroWatch v1.0 (versionCode 1, versionName '1.0')"),
        ("Mobile Frontend Framework", "Ionic Capacitor v6 + React 18 + Vite"),
        ("Android Native Shell", "Android SDK 34 (Android 14), Java 17, Gradle 9.4"),
        ("Backend Cloud API", "Django REST Framework (DRF) + Python 3.11 on Render"),
        ("Computer Vision & AI", "Ultralytics YOLOv8 (Custom Crop Models) + OpenCV + PyTorch"),
        ("Target Package Artifact", "Agrowatch_v1.0.apk")
    ]
    for idx, (lbl, val) in enumerate(tech_data):
        row = tech_table.rows[idx]
        c1, c2 = row.cells[0], row.cells[1]
        c1.width = Inches(2.5)
        c2.width = Inches(4.3)
        set_cell_background(c1, "EAEFEA")
        set_cell_background(c2, "FAFAFA")
        set_cell_margins(c1, top=80, bottom=80, left=100, right=100)
        set_cell_margins(c2, top=80, bottom=80, left=100, right=100)

        p1 = c1.paragraphs[0]
        r1 = p1.add_run(lbl)
        r1.bold = True
        r1.font.size = Pt(10)

        p2 = c2.paragraphs[0]
        r2 = p2.add_run(val)
        r2.font.size = Pt(10)

    # Save document
    doc.save(str(doc_path))
    print(f"Successfully created Word Document at: {doc_path}")

if __name__ == "__main__":
    main()
